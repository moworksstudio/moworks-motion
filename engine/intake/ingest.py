"""Ingest a PowerPoint deck for a FAITHFUL job (jobs A and C): everything needed to rebuild every slide pixel-exact.
python kit.py ingest <deck.pptx> <project name or folder> [--push-factor 1.0]
(Google Slides: File > Download > PowerPoint, or the user exports it; a local .gslides file is only a shortcut and cannot be read.)

Writes into the project (deck/: the project root keeps only the launcher, film.json, index.html and folders, 08/10/2026):
  deck/layout.json   every element of every slide in canvas px (1920 wide): geometry, crop, fill, line, rotation, text boxes with insets,
                     anchor, paragraphs, runs, font, size in px (autofit included), tracking, colour (theme colours resolved), underline
  deck/LAYOUT.md     the same, one line per element (the FAITHFUL source of truth people read)
  assets/img/        every picture once (de-duplicated by hash), original resolution
  assets/img-film/   the same pictures resampled to the size they are drawn at (Lanczos; x push-factor for photos wider than 300 px
                     that get a slow push). Chromium paints an <img> whose file size differs from its drawn size with a quality that
                     depends on seek history, so films built from originals fail kit.py verify. Originals stay untouched.
  deck/fonts.json    each font family the deck uses -> the font file copied into assets/fonts/ (found in the project, the user's and
                     Windows' font folders, the kit's Guidelines folder) or "MISSING" (ask the user for the file)
v2 (05/10/2026, copied from v1 Engine/ingest.py, same measuring): engine/intake/ingest_film.py then writes film.json, one scene per slide
rebuilt by the Library preset intake/deck-slide (the v1 SSDeck maths); `kit.py ingest <deck.pptx> <project>` does both and builds.

Limits: placeholder text whose size is only inherited from the master gets the paragraph's end size or 16 px (check the stills);
gradients are reported as "gradient", not drawn; charts, SmartArt and tables are not rebuilt (export them as pictures)."""
import sys, json, hashlib, math, os, re, shutil, colorsys
from pathlib import Path
for _s in (sys.stdout, sys.stderr):
    try: _s.reconfigure(encoding="utf-8", errors="replace")   # Windows console (cp1252) and Vietnamese slide text
    except Exception: pass

CW = 1920
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


def data_dir(proj):
    """<project>/deck/ for layout.json, LAYOUT.md, fonts.json"""
    d = Path(proj) / "deck"; d.mkdir(parents=True, exist_ok=True); return d


def data_file(proj, name):
    """deck/<name>, or <project>/<name> in a project made before 08/10/2026"""
    new, old = Path(proj) / "deck" / name, Path(proj) / name
    return new if new.exists() or not old.exists() else old


def ingest(src, proj, push_factor=1.0, say=print):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    src, proj = Path(src), Path(proj)
    p = Presentation(str(src))
    K = CW / p.slide_width; CH = round(p.slide_height * K); FPX = CW / (p.slide_width / 12700)
    img_dir = proj / "assets" / "img"; img_dir.mkdir(parents=True, exist_ok=True)
    seen, fonts, warn = {}, {}, []
    px = lambda v: None if v is None else round(v * K, 1)

    # ---- theme colours (scheme:accent1 -> #RRGGBB), with lumMod / lumOff applied ----
    theme = {}
    try:
        from lxml import etree
        mp = p.slide_masters[0].part
        for rel in mp.rels.values():
            if rel.reltype.endswith("/theme"):
                x = etree.fromstring(rel.target_part.blob)
                cs = x.find(".//a:clrScheme", NS)
                for c in (cs if cs is not None else []):
                    name = c.tag.split("}")[1]; v = c.find("a:srgbClr", NS); s = c.find("a:sysClr", NS)
                    theme[name] = (v.get("val") if v is not None else (s.get("lastClr") if s is not None else None))
    except Exception as ex: warn.append(f"theme colours not read: {ex}")

    def mods(hexv, node):
        lm = node.find("a:lumMod", NS); lo = node.find("a:lumOff", NS)
        if lm is None and lo is None: return hexv
        r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4)); h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = min(1, max(0, l * (int(lm.get("val")) / 100000 if lm is not None else 1) + (int(lo.get("val")) / 100000 if lo is not None else 0)))
        r, g, b = colorsys.hls_to_rgb(h, l, s); return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))

    def colour(el):
        if el is None: return None
        c = el.find("a:srgbClr", NS); hexv = None
        if c is not None: hexv = c.get("val")
        else:
            c = el.find("a:schemeClr", NS)
            if c is not None:
                v = theme.get(ALIAS.get(c.get("val"), c.get("val")))
                if not v: return "scheme:" + c.get("val")
                hexv = v
        if c is None or hexv is None: return None
        hexv = mods(hexv, c); a = c.find("a:alpha", NS)
        return "#" + hexv.upper() + (f"@{int(a.get('val')) / 1000:.0f}%" if a is not None else "")

    def shape_fill(sp):
        spPr = sp._element.find(".//p:spPr", NS)
        if spPr is None: spPr = sp._element.find(".//a:spPr", NS)
        if spPr is None: return None, None, None, None
        fill = None; sf = spPr.find("a:solidFill", NS)
        if sf is not None: fill = colour(sf)
        elif spPr.find("a:noFill", NS) is not None: fill = "none"
        elif spPr.find("a:gradFill", NS) is not None: fill = "gradient"
        ln = spPr.find("a:ln", NS); line = None
        if ln is not None:
            if ln.find("a:noFill", NS) is not None: line = "none"
            else:
                lf = ln.find("a:solidFill", NS); w = ln.get("w")
                line = f"{colour(lf) if lf is not None else '?'} {round(int(w) * K, 1) if w else '?'}px"
                dash = ln.find("a:prstDash", NS)
                if dash is not None and dash.get("val") != "solid": line += " " + dash.get("val")
        geom = spPr.find("a:prstGeom", NS); g = geom.get("prst") if geom is not None else ("custom" if spPr.find("a:custGeom", NS) is not None else None)
        adj = None
        if geom is not None:
            gd = geom.findall("a:avLst/a:gd", NS)
            if gd: adj = {x.get("name"): x.get("fmla") for x in gd}
        return fill, line, g, adj

    def text_info(sp, where):
        tf = sp.text_frame; body = tf._txBody.find("a:bodyPr", NS); ins = {}
        for k in ("lIns", "tIns", "rIns", "bIns"):
            v = body.get(k) if body is not None else None
            ins[k] = round(int(v) * K, 1) if v is not None else round({"lIns": 7.2, "rIns": 7.2, "tIns": 3.6, "bIns": 3.6}[k] * FPX, 1)
        anchor = body.get("anchor", "t") if body is not None else "t"; wrap = body.get("wrap", "square") if body is not None else "square"
        scale = 1.0; lnred = 0.0
        if body is not None:
            na = body.find("a:normAutofit", NS)
            if na is not None: scale = int(na.get("fontScale", "100000")) / 100000; lnred = int(na.get("lnSpcReduction", "0")) / 100000
        paras = []
        for para in tf.paragraphs:
            pPr = para._p.find("a:pPr", NS); algn = pPr.get("algn", "l") if pPr is not None else "l"; lsp = None; sb = sa = None
            if pPr is not None:
                ls = pPr.find("a:lnSpc/a:spcPct", NS)
                if ls is not None: lsp = int(ls.get("val")) / 100000
                lp = pPr.find("a:lnSpc/a:spcPts", NS)
                if lp is not None: lsp = f"{int(lp.get('val')) / 100 * FPX:.1f}px"
                b = pPr.find("a:spcBef/a:spcPts", NS); sb = round(int(b.get("val")) / 100 * FPX, 1) if b is not None else None
                b = pPr.find("a:spcAft/a:spcPts", NS); sa = round(int(b.get("val")) / 100 * FPX, 1) if b is not None else None
            runs = []; endp = para._p.find("a:endParaRPr", NS)
            end_sz = int(endp.get("sz")) / 100 if endp is not None and endp.get("sz") else None
            for child in para._p:
                tag = child.tag.split("}")[1]
                if tag == "br": runs.append({"br": True, "text": "\n"}); continue
                if tag == "fld": runs.append({"field": child.get("type"), "text": "".join(t.text or "" for t in child.findall(".//a:t", NS))}); continue
                if tag != "r": continue
                rPr = child.find("a:rPr", NS); text = "".join(t.text or "" for t in child.findall("a:t", NS))
                g = lambda k: rPr.get(k) if rPr is not None else None
                sz = int(g("sz")) / 100 if g("sz") else None
                latin = rPr.find("a:latin", NS) if rPr is not None else None; fam = latin.get("typeface") if latin is not None else None
                if fam and fam.startswith("+"): fam = None                                    # theme font reference (+mj-lt): resolved by the fallback below
                bold, ital = g("b") == "1", g("i") == "1"; spc = int(g("spc")) / 100 * FPX if g("spc") else 0
                col = colour(rPr.find("a:solidFill", NS)) if rPr is not None and rPr.find("a:solidFill", NS) is not None else None
                hl = colour(rPr.find("a:highlight", NS)) if rPr is not None and rPr.find("a:highlight", NS) is not None else None
                if text.strip():
                    fonts[(fam, bold, ital)] = fonts.get((fam, bold, ital), 0) + 1
                    if sz is None: warn.append(f"{where}: text \"{text[:30]}\" has no size of its own (inherited from the master); check the still")
                runs.append({"text": text, "font": fam, "pt": sz, "px": round(sz * FPX * scale, 1) if sz else None, "bold": bold, "italic": ital,
                             "tracking_px": round(spc * scale, 2), "cap": g("cap"), "color": col, "highlight": hl, "underline": g("u")})
            paras.append({"align": algn, "line": lsp, "before": sb, "after": sa, "runs": runs, "end_px": round(end_sz * FPX * scale, 1) if end_sz else None})
        return {"insets": ins, "anchor": anchor, "wrap": wrap, "fontScale": scale, "lnSpcReduction": lnred, "paras": paras}

    def walk(shapes, sidx, out, prefix=""):
        for z, s in enumerate(shapes):
            e = {"z": f"{prefix}{z}", "id": s.shape_id, "name": s.name, "kind": str(s.shape_type).split(" ")[0] if s.shape_type else "?",
                 "x": px(s.left), "y": px(s.top), "w": px(s.width), "h": px(s.height), "rot": round(s.rotation, 2) if hasattr(s, "rotation") else 0}
            if e["x"] is None: continue
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                e["children"] = []; out.append(e); walk(s.shapes, sidx, e["children"], prefix=f"{prefix}{z}."); continue
            if s.shape_type in (MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.TABLE) or getattr(s, "has_chart", False) or getattr(s, "has_table", False):
                warn.append(f"slide {sidx} z{e['z']}: {e['kind']} is not rebuilt; export it as a picture")
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE or (hasattr(s, "image") and s.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and _has_image(s)):
                blob = s.image.blob; h = hashlib.sha1(blob).hexdigest()
                if h not in seen:
                    fn = f"s{sidx:02d}-{s.shape_id}.{s.image.ext}"; (img_dir / fn).write_bytes(blob); seen[h] = fn
                e["img"] = "assets/img/" + seen[h]; e["img_px"] = list(s.image.size)
                e["crop"] = {k: round(getattr(s, "crop_" + k), 4) for k in ("left", "top", "right", "bottom")}
                f, l, g, adj = shape_fill(s); e["line"] = l; e["geom"] = g; e["adj"] = adj
            else:
                try:
                    f, l, g, adj = shape_fill(s); e.update({"fill": f, "line": l, "geom": g, "adj": adj})
                except Exception as ex: e["fill_err"] = str(ex)
            if getattr(s, "has_text_frame", False) and s.has_text_frame and s.text_frame.text.strip():
                e["text"] = text_info(s, f"slide {sidx} z{e['z']}")
                I = e["text"]["insets"]
                if e["h"] and I["tIns"] + I["bIns"] >= e["h"]: warn.append(f"slide {sidx} z{e['z']}: text box insets ({I['tIns']} + {I['bIns']} px) are not smaller than its height ({e['h']} px); Google Slides draws such text higher than the inset rule (credentials footer: 8 px). Check it with kit.py compare and nudge it if needed")
            out.append(e)

    def bg_of(el):
        try:
            b = el.find("p:cSld/p:bg", NS)
            if b is None: return None
            sf = b.find(".//a:solidFill", NS)
            if sf is not None: return colour(sf)
            ref = b.find("p:bgRef", NS)
            if ref is not None:
                c = colour(ref)
                if c: return c
            return "image/other"
        except Exception: return None

    slides = []
    for i, slide in enumerate(p.slides, 1):
        lay = slide.slide_layout; mas = lay.slide_master
        bg = bg_of(slide._element) or bg_of(lay._element) or bg_of(mas._element) or "#FFFFFF"
        els = []
        if slide._element.get("showMasterSp") != "0":
            if lay._element.get("showMasterSp") != "0":
                walk([s for s in mas.shapes if not s.is_placeholder], i, els, prefix="M")
            n0 = len(els); walk([s for s in lay.shapes if not s.is_placeholder], i, els, prefix="L")
            for e in els: e["from_layout"] = True
        walk(slide.shapes, i, els)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        slides.append({"slide": i, "bg": bg, "notes": notes, "elements": els})

    D = {"source": str(src), "canvas": [CW, CH], "px_per_pt": round(FPX, 4), "slides": slides}
    made = prescale(proj, D, push_factor)
    (data_dir(proj) / "layout.json").write_text(json.dumps(D, ensure_ascii=False, indent=1), encoding="utf-8")
    fmap = find_fonts(proj, fonts)
    (data_dir(proj) / "fonts.json").write_text(json.dumps(fmap, ensure_ascii=False, indent=1), encoding="utf-8")
    write_layout_md(proj, src, D, fonts, fmap, FPX)
    say(f"ingest: {len(slides)} slides, canvas {CW}x{CH}, {len(seen)} unique pictures (+{made} resampled to drawn size in assets/img-film)")
    for n in fmap.get("_notes", []): say("  font note: " + n)
    for fam, v in fmap.items():
        if fam == "_notes": continue
        say(f"  font {fam}: " + ("MISSING - ask the user for the font file, put it in assets/fonts and run ingest again" if v == "MISSING" else ", ".join(f"{x['weight']} {x['style']} <- {x['from']}" for x in v)))
    uniq, seenw = [], set()                                                    # a master/layout element repeats on every slide: say it once
    for w in warn:
        k = re.sub(r"^slide \d+ ", "", w)
        if k not in seenw: seenw.add(k); uniq.append(w)
    for w in uniq[:20]: say("  note: " + w)
    if len(uniq) > 20: say(f"  ... {len(uniq) - 20} more notes")
    return D


def _has_image(s):
    try: s.image; return True
    except Exception: return False


def prescale(proj, D, F=1.0, MINW=300):
    """resample every picture to its drawn size (x F for photos wider than MINW px that may get a slow push)"""
    from PIL import Image
    out = proj / "assets" / "img-film"; out.mkdir(parents=True, exist_ok=True); made = 0
    def walk(els):
        nonlocal made
        for e in els:
            if e.get("children"): walk(e["children"])
            if not e.get("img"): continue
            c = e.get("crop") or {"left": 0, "top": 0, "right": 0, "bottom": 0}
            iw, ih = e["w"] / max(0.01, 1 - c["left"] - c["right"]), e["h"] / max(0.01, 1 - c["top"] - c["bottom"])
            k = F if e["w"] > MINW else 1.0
            W, H = max(1, math.ceil(iw * k)), max(1, math.ceil(ih * k))
            srcf = proj / e["img"]
            try: im = Image.open(srcf)
            except Exception: continue                                              # EMF/WMF and other formats the browser cannot show either
            if (W, H) == im.size: e["img_film"] = e["img"]; e["img_film_px"] = [W, H]; continue
            name = f"{srcf.stem}_{W}x{H}{srcf.suffix}"; dst = out / name
            if not dst.exists():
                mode = "RGBA" if im.mode in ("RGBA", "LA", "P") else "RGB"
                r = im.convert(mode).resize((W, H), Image.LANCZOS)
                if srcf.suffix.lower() in (".jpg", ".jpeg"): r.convert("RGB").save(dst, quality=95, subsampling=0)
                else: r.save(dst, optimize=True)
                made += 1
            e["img_film"] = "assets/img-film/" + name; e["img_film_px"] = [W, H]
    for s in D["slides"]: walk(s["elements"])
    return made


def _font_dirs(proj):
    kit = Path(__file__).resolve().parent.parent.parent                 # v2: engine/intake/ -> the kit; brand packs carry their fonts
    sys.path.insert(0, str(kit / "engine")); from fonts.fonts import adobe_dirs
    dirs = [proj / "assets" / "fonts", Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "Windows" / "Fonts", Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts",
            kit / "brands"] + adobe_dirs()                # Adobe Fonts (Creative Cloud), numbered files without extension
    return [d for d in dirs if d.exists()]


def find_fonts(proj, used):
    """map each family used by the deck to font files (fontTools reads family / weight / italic from the file itself)"""
    want = {}
    for (fam, bold, ital) in used:
        fam = fam or "(default)"; want.setdefault(fam, set()).add((bold, ital))
    try: from fontTools.ttLib import TTFont
    except ImportError: return {f: "MISSING" for f in want}
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent)); from fonts.fonts import copy_name, font_ext
    index = {}
    for rank, d in enumerate(_font_dirs(proj)):                          # rank 0 = the project's own assets/fonts: preferred
        for f in d.rglob("*"):
            if font_ext(f) not in (".ttf", ".otf"): continue
            try:
                t = TTFont(str(f), lazy=True, fontNumber=0); nm = t["name"]
                fam = (nm.getDebugName(16) or nm.getDebugName(1) or "").strip(); sub = (nm.getDebugName(17) or nm.getDebugName(2) or "").lower()
                wt = t["OS/2"].usWeightClass if "OS/2" in t else 400; axes = {a.axisTag: a for a in t["fvar"].axes} if "fvar" in t else {}
                index.setdefault(fam.lower(), []).append({"path": f, "weight": wt, "italic": "italic" in sub or "oblique" in sub, "axes": axes, "rank": rank,
                                                          "name": copy_name(f, nm.getDebugName(6) or "")}); t.close()
            except Exception: pass
    out = {}
    fdir = proj / "assets" / "fonts"; fdir.mkdir(parents=True, exist_ok=True)
    for fam, styles in want.items():
        # v2: a brand pack's font files come from brands/<pack>/fonts like any other font (v1 loaded them through a kit stylesheet)
        cands = index.get(fam.lower(), [])
        if not cands: out[fam] = "MISSING"; continue
        faces = []
        for bold, ital in sorted(styles):
            same = [c for c in cands if c["italic"] == ital] or cands
            var = sorted((c for c in same if "wght" in c["axes"]), key=lambda c: c["rank"])
            pick = var[0] if var else min(same, key=lambda c: (abs(c["weight"] - (700 if bold else 400)), c["rank"]), default=None)
            if not pick: continue
            if not var and pick["weight"] != (700 if bold else 400): out.setdefault("_notes", []).append(f"{fam} {'bold' if bold else 'regular'}{' italic' if ital else ''}: no exact file, using weight {pick['weight']} ({pick['path'].name}); ask the user for the right file")
            if ital and not pick["italic"]: out.setdefault("_notes", []).append(f"{fam} italic: no italic file, the browser will slant the upright one")
            dst = fdir / pick["name"]
            if pick["path"].resolve() != dst.resolve() and not dst.exists(): shutil.copy2(pick["path"], dst)
            if "wght" in pick["axes"]: a = pick["axes"]["wght"]; weight = f"{int(a.minValue)} {int(a.maxValue)}"
            else: weight = str(pick["weight"])
            face = {"src": "assets/fonts/" + dst.name, "weight": weight, "style": "italic" if pick["italic"] else "normal", "from": str(pick["path"])}
            if "opsz" in pick["axes"]: face["opsz"] = "none"               # variable fonts with optical size (Inter 4): slide apps draw the default cut
            if face not in faces: faces.append(face)
        out[fam] = faces or "MISSING"
    return out


def write_layout_md(proj, src, D, fonts, fmap, FPX):
    slides = D["slides"]; CW, CH = D["canvas"]
    L = ["# LAYOUT (FAITHFUL source of truth, written by kit.py ingest)", "",
         f"Source: `{src}` ({len(slides)} slides). Canvas {CW}x{CH} px. 1 pt = {FPX:.4f} px. Positions are top-left x, y and w x h in canvas px; z = paint order (0 = bottom).",
         "Font px already include the text box's autofit scale. Full data (insets, line spacing, crops): `layout.json`. zM* / zL* = master / layout elements (painted under the slide's own shapes). FIELD slidenum = slide number.", ""]
    for s in slides:
        L.append(f"## Slide {s['slide']:02d}  (bg {s['bg']})")
        if s["notes"]: L.append(f"Notes: {s['notes'][:300]}")
        def line(e, ind=""):
            base = f"{ind}- z{e['z']} {e['kind'].lower()} @ {e['x']},{e['y']} {e['w']}x{e['h']}" + (f" rot{e['rot']}" if e.get("rot") else "")
            if "img" in e:
                c = e["crop"]; cs = "" if not any(c.values()) else f" crop l{c['left']} t{c['top']} r{c['right']} b{c['bottom']}"
                base += f" IMG `{e['img']}` ({e['img_px'][0]}x{e['img_px'][1]}){cs}" + (f" geom {e['geom']}" if e.get("geom") not in (None, "rect") else "")
            else:
                bits = [f"fill {e['fill']}"] if e.get("fill") not in (None, "none") else []
                if e.get("line") not in (None, "none"): bits.append(f"line {e['line']}")
                if e.get("geom") not in (None, "rect"): bits.append(f"geom {e['geom']}" + (f" {e['adj']}" if e.get("adj") else ""))
                if bits: base += " " + ", ".join(bits)
            if "text" in e:
                t = e["text"]
                for pa in t["paras"]:
                    for r in pa["runs"]:
                        if "field" in r: base += f" | FIELD {r['field']}"; continue
                        if r.get("br"): base += " | <br>"; continue
                        if not r["text"].strip(): continue
                        st = f"{r['font']} {r['px']}px" + (" B" if r["bold"] else "") + (" I" if r["italic"] else "") + (f" trk{r['tracking_px']}" if r["tracking_px"] else "") + (f" {r['color']}" if r["color"] else "") + (f" cap:{r['cap']}" if r["cap"] and r["cap"] != "none" else "")
                        base += f" | \"{r['text']}\" [{st}, {pa['align']}" + (f", lh {pa['line']}" if pa["line"] else "") + "]"
                if t["fontScale"] != 1: base += f" (autofit {t['fontScale']})"
                if t["anchor"] != "t": base += f" (anchor {t['anchor']})"
            L.append(base)
            for c in e.get("children", []): line(c, ind + "  ")
        for e in s["elements"]: line(e)
        L.append("")
    L.append("## Fonts used (family, bold, italic: run count) and the files found")
    for (f, b, it), n in sorted(fonts.items(), key=lambda x: -x[1]): L.append(f"- {f} {'Bold ' if b else ''}{'Italic' if it else ''}: {n}")
    for fam, v in fmap.items():
        if fam == "_notes": L += [f"- note: {n}" for n in v]; continue
        L.append(f"- file for {fam}: " + ("MISSING" if v == "MISSING" else ", ".join(x.get("src", x.get("from")) for x in v)))
    (data_dir(proj) / "LAYOUT.md").write_text("\n".join(L) + "\n", encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) < 3: raise SystemExit(__doc__)
    pf = float(sys.argv[sys.argv.index("--push-factor") + 1]) if "--push-factor" in sys.argv else 1.0
    ingest(sys.argv[1], sys.argv[2], pf)
